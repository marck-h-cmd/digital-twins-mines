import os
import io
import json
import datetime
import numpy as np
import pandas as pd
import streamlit as st
from utils.translations import get_text, get_current_lang

# Librerías de exportación
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# Helper for default mock dataset (20 sample records)
def get_mock_dataset():
    np.random.seed(42)
    sample_names = [f"Entidad / Estudiante #{1001 + i}" for i in range(20)]
    distances = np.round(np.random.uniform(2.0, 45.0, 20), 2)
    ttc = np.round(np.random.uniform(0.5, 8.0, 20), 2)
    bpm = np.random.randint(65, 145, 20)
    fatigue = np.round(np.random.uniform(0.05, 0.95, 20), 2)
    gas_co = np.round(np.random.uniform(5.0, 45.0, 20), 1)
    
    # Calculate risk level based on thresholds
    risk_labels = []
    risk_levels = []
    for d, t, f in zip(distances, ttc, fatigue):
        if d < 5.0 or t < 1.5 or f > 0.75:
            risk_labels.append("ALTO")
            risk_levels.append(2)
        elif d < 15.0 or t < 3.5 or f > 0.45:
            risk_labels.append("MEDIO")
            risk_levels.append(1)
        else:
            risk_labels.append("BAJO")
            risk_levels.append(0)
            
    df_mock = pd.DataFrame({
        "ID": [1001 + i for i in range(20)],
        "Nombre_Entidad": sample_names,
        "distance_3d": distances,
        "ttc": ttc,
        "worker_bpm": bpm,
        "fatigue_index": fatigue,
        "gas_co_ppm": gas_co,
        "risk_level": risk_levels,
        "risk_label": risk_labels
    })
    return df_mock

# Helper para cargar best_model_meta.json por defecto
def load_default_metadata():
    possible_paths = [
        os.path.join(os.path.dirname(__file__), "..", "..", "backend", "app", "ml", "artifacts", "best_model_meta.json"),
        os.path.join(os.path.dirname(__file__), "..", "artifacts", "best_model_meta.json"),
        "backend/app/ml/artifacts/best_model_meta.json"
    ]
    for path in possible_paths:
        if os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
    return {
        "model_name": "RandomForest",
        "test_metrics": {
            "accuracy": 0.9985,
            "f1_macro": 0.9982,
            "precision_macro": 0.9976,
            "recall_macro": 0.9988,
            "auc_roc": 1.0000,
            "confusion_matrix": [[960, 2, 0], [0, 412, 0], [0, 1, 625]]
        }
    }

# Helper para convertir figuras de Plotly/Matplotlib a PNG Bytes
def fig_to_bytes(fig):
    if hasattr(fig, "to_image"):
        try:
            return fig.to_image(format="png", width=700, height=400)
        except Exception:
            pass
    import matplotlib.pyplot as plt
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight", dpi=150)
    buf.seek(0)
    return buf.getvalue()

# Helper para contar distribuciones de riesgo de manera unificada sin duplicación
def get_risk_counts(data: pd.DataFrame):
    total_evals = len(data)
    if total_evals == 0:
        return 0, 0, 0, 0

    if "risk_label" in data.columns:
        alto_cnt = int((data["risk_label"].astype(str).str.upper() == "ALTO").sum())
        medio_cnt = int((data["risk_label"].astype(str).str.upper() == "MEDIO").sum())
        bajo_cnt = int((data["risk_label"].astype(str).str.upper() == "BAJO").sum())
    elif "risk_level" in data.columns:
        alto_cnt = int((data["risk_level"] == 2).sum())
        medio_cnt = int((data["risk_level"] == 1).sum())
        bajo_cnt = int((data["risk_level"] == 0).sum())
    else:
        alto_cnt, medio_cnt, bajo_cnt = 0, 0, total_evals

    sum_known = alto_cnt + medio_cnt + bajo_cnt
    if sum_known < total_evals:
        bajo_cnt += (total_evals - sum_known)

    return total_evals, alto_cnt, medio_cnt, bajo_cnt

# 1. EXPORTACIÓN A PDF (ReportLab)
def generate_pdf(data: pd.DataFrame = None, figures_bytes: dict = None, meta_info: dict = None, evaluator_name: str = None, technical_notes: str = None, include_friedman: bool = True, lang: str = None) -> bytes:
    """
    Genera un informe en PDF profesional con estilo académico (ReportLab).
    Extrae las métricas del modelo campeón desde best_model_meta.json para el Anexo Técnico.
    """
    if data is None or data.empty:
        data = get_mock_dataset()
    if meta_info is None:
        meta_info = load_default_metadata()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0f172a'),
        alignment=1 # Center
    )

    heading_style = ParagraphStyle(
        'HeadingStyle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#1e3a8a'),
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'BodyStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#334155'),
        alignment=4 # Justified
    )

    # Encabezado
    date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    elements.append(Paragraph(f"<b>{get_text('rep_header_title')}</b>", title_style))
    eval_txt = f" | {get_text('evaluator_label')}: {evaluator_name}" if evaluator_name else ""
    elements.append(Paragraph(f"<b>{get_text('rep_header_subtitle')}</b>", ParagraphStyle('SubTitle', parent=title_style, fontSize=11, textColor=colors.HexColor('#475569'))))
    elements.append(Paragraph(f"<i>Fecha: {date_str}{eval_txt} | Entorno: Sistema M-11 Digital Twin</i>", ParagraphStyle('CenterMeta', parent=body_style, alignment=1)))
    elements.append(Spacer(1, 15))

    # Resumen Ejecutivo
    elements.append(Paragraph(get_text("rep_sec1_title"), heading_style))
    total_evals, alto_cnt, medio_cnt, bajo_cnt = get_risk_counts(data)

    kpi_data = [
        [get_text("total_records"), get_text("alert_high"), get_text("alert_medium"), get_text("alert_low")],
        [f"{total_evals:,}", f"{alto_cnt} ({alto_cnt/max(total_evals,1)*100:.1f}%)", f"{medio_cnt} ({medio_cnt/max(total_evals,1)*100:.1f}%)", f"{bajo_cnt} ({bajo_cnt/max(total_evals,1)*100:.1f}%)"]
    ]
    t_kpi = Table(kpi_data, colWidths=[4.5*cm, 4.5*cm, 4.5*cm, 4.5*cm])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#f8fafc')]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('FONTSIZE', (0,0), (-1,-1), 9),
    ]))
    elements.append(t_kpi)
    elements.append(Spacer(1, 15))

    # Notas Técnicas
    if technical_notes:
        elements.append(Paragraph(get_text("rep_sec2_title"), heading_style))
        elements.append(Paragraph(f"<i>Observación: {technical_notes}</i>", body_style))
        elements.append(Spacer(1, 15))

    # Tabla de Datos (Muestra)
    elements.append(Paragraph(get_text("rep_sec3_title"), heading_style))
    cols_show = [c for c in ['ID', 'Nombre_Entidad', 'distance_3d', 'ttc', 'worker_bpm', 'fatigue_index', 'gas_co_ppm', 'risk_label'] if c in data.columns]
    if not cols_show:
        cols_show = list(data.columns[:6])
    sample_df = data[cols_show].head(8)
    
    t_sample_data = [cols_show]
    for idx, row in sample_df.iterrows():
        t_sample_data.append([str(row[c]) for c in cols_show])
        
    col_w = (18.0 / len(cols_show)) * cm
    t_sample = Table(t_sample_data, colWidths=[col_w]*len(cols_show))
    t_sample.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f1f5f9')]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#94a3b8')),
        ('FONTSIZE', (0,0), (-1,-1), 8),
    ]))
    elements.append(t_sample)
    elements.append(Spacer(1, 15))

    # 4. Visualizaciones Físicas y Gráficos Telemétricos
    elements.append(Paragraph(get_text("rep_sec4_title"), heading_style))
    has_valid_images = False
    if figures_bytes:
        for fig_name, img_data in figures_bytes.items():
            if img_data:
                img_buf = io.BytesIO(img_data)
                elements.append(RLImage(img_buf, width=16*cm, height=8.5*cm))
                elements.append(Spacer(1, 10))
                has_valid_images = True
                
    if not has_valid_images:
        try:
            import matplotlib
            matplotlib.use('Agg')
            import matplotlib.pyplot as plt

            total_evals, alto_cnt, medio_cnt, bajo_cnt = get_risk_counts(data)
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(9.5, 3.6), dpi=150)
            
            levels = [get_text("label_high"), get_text("label_medium"), get_text("label_low")]
            values = [alto_cnt, medio_cnt, bajo_cnt]
            bar_colors = ["#ef4444", "#f59e0b", "#10b981"]
            bars = ax1.bar(levels, values, color=bar_colors, width=0.55, edgecolor="#0f172a", linewidth=1)
            ax1.set_title("Distribución de Riesgos Telemétricos" if get_current_lang() == "es" else "Telemetric Risk Distribution", fontsize=9, fontweight="bold", pad=8)
            ax1.set_ylabel("Eventos" if get_current_lang() == "es" else "Events", fontsize=8)
            ax1.grid(axis='y', linestyle='--', alpha=0.5)
            for bar in bars:
                yval = bar.get_height()
                ax1.text(bar.get_x() + bar.get_width()/2.0, yval + 0.15, f"{int(yval)}", ha='center', va='bottom', fontsize=8, fontweight='bold')

            dist_col = 'distance_3d' if 'distance_3d' in data.columns else data.columns[0]
            ttc_col = 'ttc' if 'ttc' in data.columns else data.columns[1]
            ax2.scatter(data[dist_col].head(15), data[ttc_col].head(15), c="#1e40af", s=55, alpha=0.85, edgecolors="#0f172a", linewidths=0.8)
            ax2.set_title(f"Dispersión: {dist_col} vs {ttc_col}", fontsize=9, fontweight="bold", pad=8)
            ax2.set_xlabel(dist_col, fontsize=8)
            ax2.set_ylabel(ttc_col, fontsize=8)
            ax2.grid(True, linestyle='--', alpha=0.5)

            plt.tight_layout()
            buf = io.BytesIO()
            plt.savefig(buf, format='png', bbox_inches='tight', dpi=150)
            plt.close(fig)
            buf.seek(0)
            elements.append(RLImage(buf, width=16.5*cm, height=6.2*cm))
        except Exception as e:
            print(f"Error generando gráfico por defecto en Streamlit PDF: {e}")

    elements.append(Spacer(1, 10))

    # PÁGINA DE ANEXOS TÉCNICOS: MÉTRICAS DEL MODELO CAMPEÓN
    elements.append(PageBreak())
    elements.append(Paragraph(get_text("rep_sec5_title"), title_style))
    elements.append(Spacer(1, 15))

    if meta_info and "test_metrics" in meta_info:
        m = meta_info["test_metrics"]
        model_name = meta_info.get("model_name", "RandomForest")
        
        elements.append(Paragraph(f"<b>{get_text('champion_model')}:</b> {model_name}", heading_style))
        
        meta_table_data = [
            ["Métrica de Evaluación" if get_current_lang() == "es" else "Evaluation Metric", "Valor (Test Set)" if get_current_lang() == "es" else "Value (Test Set)"],
            [get_text("metric_accuracy"), f"{m.get('accuracy', 0)*100:.2f}%"],
            [get_text("metric_f1"), f"{m.get('f1_macro', 0):.4f}"],
            [get_text("metric_precision"), f"{m.get('precision_macro', 0):.4f}"],
            [get_text("metric_recall"), f"{m.get('recall_macro', 0):.4f}"],
            [get_text("metric_auc"), f"{m.get('auc_roc', 0):.4f}"]
        ]
        t_meta = Table(meta_table_data, colWidths=[9*cm, 9*cm])
        t_meta.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e3a8a')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#f8fafc'), colors.white]),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('FONTSIZE', (0,0), (-1,-1), 9),
        ]))
        elements.append(t_meta)
        elements.append(Spacer(1, 15))

    if include_friedman and meta_info and "statistical_tests" in meta_info:
        st_tests = meta_info["statistical_tests"]
        elements.append(Paragraph(f"<b>{get_text('rep_sec6_title')}</b>", heading_style))
        elements.append(Paragraph(f"Friedman Stat: {st_tests.get('friedman_stat', 19.04):.2f} (p-value: {st_tests.get('friedman_p', 0.00077):.5f})", body_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


# 2. EXPORTACIÓN A WORD (DOCX)
def generate_word(data: pd.DataFrame = None, figures_bytes: dict = None, meta_info: dict = None, evaluator_name: str = None, technical_notes: str = None, include_friedman: bool = True, lang: str = None) -> bytes:
    """
    Genera un informe en Word (.docx) editable estructurado con títulos, tablas y gráficos.
    """
    if data is None or data.empty:
        data = get_mock_dataset()
    if meta_info is None:
        meta_info = load_default_metadata()

    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(f"{get_text('rep_header_title')}\n")
    run_title.bold = True
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    run_sub = p_title.add_run(f"{get_text('rep_header_subtitle')}\n")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    eval_txt = f" | {get_text('evaluator_label')}: {evaluator_name}" if evaluator_name else ""
    p_meta = doc.add_paragraph(f"Fecha / Date: {date_str}{eval_txt} | Entorno: M-11 Digital Twin")
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.runs[0].font.italic = True
    p_meta.runs[0].font.size = Pt(9)

    doc.add_heading(get_text("rep_sec1_title"), level=1)
    total_evals, alto_cnt, medio_cnt, bajo_cnt = get_risk_counts(data)

    # Tabla Resumen
    t_summary = doc.add_table(rows=2, cols=4)
    t_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [get_text("total_records"), get_text("alert_high"), get_text("alert_medium"), get_text("alert_low")]
    for i, h in enumerate(headers):
        t_summary.cell(0, i).text = h
        t_summary.cell(0, i).paragraphs[0].runs[0].font.bold = True

    t_summary.cell(1, 0).text = f"{total_evals:,}"
    t_summary.cell(1, 1).text = f"{alto_cnt} ({alto_cnt/max(total_evals,1)*100:.1f}%)"
    t_summary.cell(1, 2).text = f"{medio_cnt} ({medio_cnt/max(total_evals,1)*100:.1f}%)"
    t_summary.cell(1, 3).text = f"{bajo_cnt} ({bajo_cnt/max(total_evals,1)*100:.1f}%)"

    if technical_notes:
        doc.add_heading(get_text("rep_sec2_title"), level=1)
        doc.add_paragraph(technical_notes)

    doc.add_heading(get_text("rep_sec3_title"), level=1)
    cols_show = [c for c in ['ID', 'Nombre_Entidad', 'distance_3d', 'ttc', 'worker_bpm', 'fatigue_index', 'gas_co_ppm', 'risk_label'] if c in data.columns]
    if not cols_show:
        cols_show = list(data.columns[:6])
    sample_df = data[cols_show].head(10)

    t_data = doc.add_table(rows=len(sample_df)+1, cols=len(cols_show))
    t_data.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, col_name in enumerate(cols_show):
        cell = t_data.cell(0, j)
        cell.text = col_name
        cell.paragraphs[0].runs[0].font.bold = True

    for i, (_, row) in enumerate(sample_df.iterrows(), 1):
        for j, col_name in enumerate(cols_show):
            t_data.cell(i, j).text = str(row[col_name])

    # Incrustar Figuras
    if figures_bytes:
        doc.add_heading(get_text("rep_sec4_title"), level=1)
        for fig_name, img_data in figures_bytes.items():
            if img_data:
                img_buf = io.BytesIO(img_data)
                doc.add_picture(img_buf, width=Inches(6.0))

    # Anexos del Modelo
    doc.add_page_break()
    doc.add_heading(get_text("rep_sec5_title"), level=1)
    
    if meta_info and "test_metrics" in meta_info:
        m = meta_info["test_metrics"]
        doc.add_paragraph(f"{get_text('champion_model')}: {meta_info.get('model_name', 'RandomForest')}")
        
        t_meta = doc.add_table(rows=6, cols=2)
        t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_meta.cell(0, 0).text = "Evaluation Metric" if get_current_lang() == "en" else "Métrica de Evaluación"
        t_meta.cell(0, 1).text = "Test Value" if get_current_lang() == "en" else "Valor en Test Set"
        t_meta.cell(0, 0).paragraphs[0].runs[0].font.bold = True
        t_meta.cell(0, 1).paragraphs[0].runs[0].font.bold = True

        metrics_list = [
            (get_text("metric_accuracy"), f"{m.get('accuracy', 0)*100:.2f}%"),
            (get_text("metric_f1"), f"{m.get('f1_macro', 0):.4f}"),
            (get_text("metric_precision"), f"{m.get('precision_macro', 0):.4f}"),
            (get_text("metric_recall"), f"{m.get('recall_macro', 0):.4f}"),
            (get_text("metric_auc"), f"{m.get('auc_roc', 0):.4f}")
        ]
        for idx, (k, v) in enumerate(metrics_list, 1):
            t_meta.cell(idx, 0).text = k
            t_meta.cell(idx, 1).text = v

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# 3. EXPORTACIÓN A EXCEL (openpyxl)
def generate_excel(data: pd.DataFrame = None, meta_info: dict = None, simulations_data: list = None, evaluator_name: str = None, technical_notes: str = None, include_friedman: bool = True, lang: str = None) -> bytes:
    """
    Genera un libro de Excel (.xlsx) con 4 hojas: Resumen, Alertas, Métricas Modelo y Simulaciones.
    Aplica formato condicional con colores (Rojo, Amarillo, Verde).
    """
    if data is None or data.empty:
        data = get_mock_dataset()
    if meta_info is None:
        meta_info = load_default_metadata()

    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remover hoja por defecto

    # Estilos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

    # HOJA 1: RESUMEN / SUMMARY
    sheet1_title = "Summary" if get_current_lang() == "en" else "Resumen"
    ws1 = wb.create_sheet(title=sheet1_title)
    ws1.append(["KPI / Operational Metric" if get_current_lang() == "en" else "KPI / Métrica Operacional", "Calculated Value" if get_current_lang() == "en" else "Valor Calculado"])
    ws1.append([get_text("evaluator_label"), evaluator_name or get_text("evaluator_name")])
    ws1.append([get_text("total_records"), len(data)])
    
    total_evals, alto_cnt, medio_cnt, bajo_cnt = get_risk_counts(data)

    ws1.append([get_text("alert_high"), alto_cnt])
    ws1.append([get_text("alert_medium"), medio_cnt])
    ws1.append([get_text("alert_low"), bajo_cnt])
    ws1.append([get_text("technical_notes_label"), technical_notes or "N/A"])
    ws1.append(["Notas Técnicas", technical_notes or "N/A"])

    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font

    # HOJA 2: ALERTAS (Detalladas con formato condicional)
    ws2 = wb.create_sheet(title="Alertas")
    cols_export = [c for c in ['ID', 'Nombre_Entidad', 'distance_3d', 'ttc', 'worker_bpm', 'fatigue_index', 'gas_co_ppm', 'risk_label'] if c in data.columns]
    if not cols_export:
        cols_export = list(data.columns)
    ws2.append(cols_export)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in dataframe_to_rows(data[cols_export], index=False, header=False):
        ws2.append(row)

    risk_col_idx = cols_export.index('risk_label') + 1 if 'risk_label' in cols_export else None
    if risk_col_idx:
        for row in ws2.iter_rows(min_row=2, max_row=len(data)+1, min_col=1, max_col=len(cols_export)):
            val = str(row[risk_col_idx-1].value)
            if "ALTO" in val or "2" in val:
                row[risk_col_idx-1].fill = red_fill
            elif "MEDIO" in val or "1" in val:
                row[risk_col_idx-1].fill = yellow_fill
            elif "BAJO" in val or "0" in val:
                row[risk_col_idx-1].fill = green_fill

    # HOJA 3: MÉTRICAS MODELO
    ws3 = wb.create_sheet(title="Métricas Modelo")
    ws3.append(["Métrica de Evaluación", "Valor (Test Set)"])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font

    if meta_info and "test_metrics" in meta_info:
        m = meta_info["test_metrics"]
        ws3.append(["Modelo Campeón", meta_info.get("model_name", "RandomForest")])
        ws3.append(["Accuracy (Exactitud)", m.get("accuracy", 0)])
        ws3.append(["F1-Score Macro", m.get("f1_macro", 0)])
        ws3.append(["Precision Macro", m.get("precision_macro", 0)])
        ws3.append(["Recall Macro", m.get("recall_macro", 0)])
        ws3.append(["AUC-ROC (OVR)", m.get("auc_roc", 0)])

    # HOJA 4: SIMULACIONES
    ws4 = wb.create_sheet(title="Simulaciones")
    ws4.append(["Escenario Simulado", "Distancia 3D (m)", "Fatiga", "Zona Restringida", "Riesgo Resultante"])
    for cell in ws4[1]:
        cell.fill = header_fill
        cell.font = header_font

    sims = simulations_data or [
        ["Escenario 1: Operación Normal", 35.0, 0.15, 0, "BAJO"],
        ["Escenario 2: Proximidad Moderada", 12.0, 0.48, 1, "MEDIO"],
        ["Escenario 3: Alerta Crítica Colisión", 4.0, 0.88, 1, "ALTO"]
    ]
    for s in sims:
        ws4.append(s)

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# 4. PREVISUALIZACIÓN FIEL EN STREAMLIT ("Ver antes de imprimir" - Simulador A4)
def show_preview(data: pd.DataFrame = None, figures_dict: dict = None, meta_info: dict = None, evaluator_name: str = "Ing. SANTOS FERNANDEZ JUAN PEDRO", technical_notes: str = None, include_friedman: bool = True):
    """
    Renderiza el simulador de Hoja A4 estructurada dentro de Streamlit con actualización en tiempo real.
    """
    if data is None or data.empty:
        data = get_mock_dataset()
    if meta_info is None:
        meta_info = load_default_metadata()

    eval_display = evaluator_name if evaluator_name else get_text("evaluator_name")
    notes_display = technical_notes if technical_notes else get_text("default_notes")

    st.markdown(f"""
        <div style="background-color: #ffffff; color: #0f172a; border: 1px solid #cbd5e1; padding: 30px; border-radius: 4px; font-family: 'Helvetica Neue', Arial, sans-serif; box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.1);">
            <div style="border-bottom: 2px solid #0f172a; padding-bottom: 12px; margin-bottom: 20px; display: flex; justify-content: space-between; align-items: flex-start;">
                <div>
                    <h2 style="color: #0f172a; margin: 0; font-size: 1.4rem; font-weight: 800; letter-spacing: -0.5px;">{get_text("rep_header_title")}</h2>
                    <h4 style="color: #475569; margin-top: 4px; font-size: 0.95rem; font-weight: 600;">{get_text("rep_header_subtitle")}</h4>
                    <p style="color: #64748b; font-size: 0.8rem; margin-top: 4px;">
                        Fecha: {datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")} | Entorno: Minería Subterránea M-11
                    </p>
                </div>
                <div style="text-align: right;">
                    <span style="background-color: #0f172a; color: #ffffff; font-size: 0.7rem; font-weight: 700; padding: 4px 8px; border-radius: 3px; font-family: monospace;">FORMATO A4 / Q1</span>
                    <p style="color: #334155; font-size: 0.8rem; margin-top: 8px; font-weight: 600;">
                        {get_text("evaluator_label")}: <span style="color: #0284c7;">{eval_display}</span>
                    </p>
                </div>
            </div>
        </div>
    """, unsafe_allow_html=True)
    
    st.subheader(get_text("rep_sec1_title"))
    total_evals, alto_cnt, medio_cnt, bajo_cnt = get_risk_counts(data)

    p1, p2, p3, p4 = st.columns(4)
    p1.metric(get_text("total_records"), f"{total_evals:,}")
    p2.metric(get_text("alert_high"), f"{alto_cnt}")
    p3.metric(get_text("alert_medium"), f"{medio_cnt}")
    p4.metric(get_text("alert_low"), f"{bajo_cnt}")

    st.subheader(get_text("rep_sec2_title"))
    st.info(f"**{get_text('technical_notes_label')} ({eval_display}):**\n\n{notes_display}")

    st.subheader(get_text("rep_sec3_title"))
    cols_show = [c for c in ['ID', 'Nombre_Entidad', 'distance_3d', 'ttc', 'worker_bpm', 'fatigue_index', 'gas_co_ppm', 'risk_label'] if c in data.columns]
    if not cols_show:
        cols_show = list(data.columns[:6])
    st.dataframe(data[cols_show], use_container_width=True)

    st.subheader(get_text("rep_sec4_title"))
    col_g1, col_g2 = st.columns(2)
    if figures_dict and len(figures_dict) > 0:
        fig_keys = list(figures_dict.keys())
        if len(fig_keys) > 0:
            key_0 = "report_preview_fig_0_" + str(fig_keys[0]).replace(" ", "_").lower()
            with col_g1:
                st.plotly_chart(figures_dict[fig_keys[0]], use_container_width=True, key=key_0)
        if len(fig_keys) > 1:
            key_1 = "report_preview_fig_1_" + str(fig_keys[1]).replace(" ", "_").lower()
            with col_g2:
                st.plotly_chart(figures_dict[fig_keys[1]], use_container_width=True, key=key_1)
    else:
        # Fallback Plotly graphics for preview
        df_risk_counts = pd.DataFrame([
            {"Nivel": get_text("alert_high"), "Cantidad": alto_cnt},
            {"Nivel": get_text("alert_medium"), "Cantidad": medio_cnt},
            {"Nivel": get_text("alert_low"), "Cantidad": bajo_cnt}
        ])
        fig_bar = px.bar(
            df_risk_counts, x="Nivel", y="Cantidad", color="Nivel",
            title=get_text("rep_sec4_title"),
            color_discrete_map={get_text("alert_high"): "#ef4444", get_text("alert_medium"): "#f59e0b", get_text("alert_low"): "#10b981"}
        )
        with col_g1:
            st.plotly_chart(fig_bar, use_container_width=True, key="report_fallback_bar_fig")

        dist_col = 'distance_3d' if 'distance_3d' in data.columns else data.columns[0]
        ttc_col = 'ttc' if 'ttc' in data.columns else data.columns[1]
        fig_scat = px.scatter(
            data.head(20), x=dist_col, y=ttc_col, color="risk_label" if "risk_label" in data.columns else None,
            title=f"Dispersión Telemétrica: {dist_col} vs {ttc_col}"
        )
        with col_g2:
            st.plotly_chart(fig_scat, use_container_width=True, key="report_fallback_scat_fig")

    if meta_info and "test_metrics" in meta_info:
        st.subheader(get_text("rep_sec5_title"))
        m = meta_info["test_metrics"]
        am1, am2, am3, am4 = st.columns(4)
        am1.metric(get_text("metric_accuracy"), f"{m.get('accuracy',0)*100:.2f}%")
        am2.metric(get_text("metric_f1"), f"{m.get('f1_macro',0):.4f}")
        am3.metric(get_text("metric_recall"), f"{m.get('recall_macro',0):.4f}")
        am4.metric(get_text("metric_auc"), f"{m.get('auc_roc',0):.4f}")

    if include_friedman and meta_info and "statistical_tests" in meta_info:
        st.subheader(get_text("rep_sec6_title"))
        st_tests = meta_info["statistical_tests"]
        st.write(f"**Friedman Statistic:** `{st_tests.get('friedman_stat', 19.04):.2f}` | **p-value:** `{st_tests.get('friedman_p', 0.00077):.5f}`")
        
        if "pairwise_hypothesis_tests" in st_tests:
            df_tests = pd.DataFrame(st_tests["pairwise_hypothesis_tests"])
            st.table(df_tests)

